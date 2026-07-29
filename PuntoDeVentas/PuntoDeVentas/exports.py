"""Utilidades genericas para exportar listados a Excel y PDF.

Cualquier vista de cualquier app puede usarlas: solo arma sus propios
encabezados y filas (ya como texto/numeros listos para mostrar) y le
pasa el trabajo de generar el archivo a estas funciones. Evita repetir
la logica de armado de Workbook/PDF en cada app.
"""
from datetime import datetime

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLOR_MARCA = "2F6FED"


def exportar_excel(nombre_archivo, encabezados, filas, titulo_hoja="Datos"):
    """Genera un .xlsx y lo devuelve como HttpResponse listo para descargar.

    encabezados: lista de strings.
    filas: lista de listas, mismo orden/cantidad que encabezados. Los
    valores se insertan tal cual (str, int, float, date), openpyxl los
    tipa automaticamente en la celda.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo_hoja[:31]  # Excel limita el nombre de hoja a 31 caracteres

    ws.append(list(encabezados))
    for col_num in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col_num)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color=COLOR_MARCA, end_color=COLOR_MARCA, fill_type="solid")

    for fila in filas:
        ws.append(list(fila))

    for col_num, encabezado in enumerate(encabezados, start=1):
        valores_columna = [str(fila[col_num - 1]) for fila in filas]
        ancho = max([len(str(encabezado))] + [len(v) for v in valores_columna] + [8])
        ws.column_dimensions[get_column_letter(col_num)].width = min(ancho + 4, 40)

    ws.freeze_panes = "A2"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}.xlsx"'
    wb.save(response)
    return response


def exportar_pdf(nombre_archivo, titulo, encabezados, filas, subtitulo=None):
    """Genera un PDF con una tabla (orientacion horizontal, para listados
    con varias columnas) y lo devuelve como HttpResponse."""
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=landscape(letter),
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
    )
    estilos = getSampleStyleSheet()
    elementos = [Paragraph(titulo, estilos["Title"])]

    pie = subtitulo or f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    elementos.append(Paragraph(pie, estilos["Normal"]))
    elementos.append(Spacer(1, 12))

    datos_tabla = [list(encabezados)] + [list(fila) for fila in filas]
    tabla = Table(datos_tabla, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{COLOR_MARCA}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F9")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    elementos.append(tabla)

    if not filas:
        elementos.append(Spacer(1, 12))
        elementos.append(Paragraph("No hay datos para este filtro.", estilos["Normal"]))

    doc.build(elementos)
    return response
